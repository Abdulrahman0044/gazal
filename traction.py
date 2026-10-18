"""Account-scoped traction ingestion. Imported snapshots never add to existing values."""
import csv,io,json,math,secrets,hashlib,hmac,time
from decimal import Decimal
from datetime import date
from fastapi import HTTPException,Depends,Request,UploadFile,File
from pydantic import BaseModel,Field
from openpyxl import load_workbook
METRICS={'total_users':('Total users','count','snapshot'),'active_users':('Active users','count','snapshot'),'paying_customers':('Paying customers','count','snapshot'),'billed':('Billed','NGN','flow'),'collected':('Collected','NGN','flow'),'refunds':('Refunds','NGN','flow'),'mrr':('MRR','NGN','snapshot')}
def install(app,conn,current,enforce,uid,now):
    # Schema is initialized centrally before route registration.
    def scope(c,u,a,p,write=False):
        enforce(c,u,a,{'account_admin','project_manager'} if write else {'account_admin','project_manager','team_lead'})
        if not c.execute('SELECT 1 FROM projects WHERE id=? AND account_id=?',(p,a)).fetchone():raise HTTPException(404,'Project not found')
    def valid(rows):
        out=[];keys=set()
        for i,r in enumerate(rows,2):
            try:
                metric=str(r['metric']);day=date.fromisoformat(str(r['date'])).isoformat();v=float(r['value'])
                if metric not in METRICS or not math.isfinite(v) or v<0 or v>1_000_000_000_000 or (METRICS[metric][1]=='count' and not v.is_integer()):raise ValueError()
                if METRICS[metric][1]=='NGN' and Decimal(str(v))*100!=(Decimal(str(v))*100).to_integral_value():raise ValueError()
                if (metric,day) in keys:raise ValueError()
                keys.add((metric,day));out.append({'metric':metric,'date':day,'value':v})
            except (ValueError,KeyError,TypeError):raise HTTPException(422,f'Invalid or duplicate row {i}; use date, metric, value with nonnegative values and whole-number counts')
        if not out or len(out)>10000:raise HTTPException(422,'Import must contain 1–10,000 rows')
        return out
    def put(c,a,p,row,s):
        c.execute('INSERT INTO traction_values VALUES(?,?,?,?,?,?,?) ON CONFLICT(account_id,project_id,metric,day) DO UPDATE SET value=excluded.value,source_id=excluded.source_id,updated_at=excluded.updated_at',(a,p,row['metric'],row['date'],row['value'],s,now()))
    class Source(BaseModel):
        metric:str;mode:str
    @app.get('/api/traction/definitions')
    def definitions(u=Depends(current)):return [{'key':k,'label':v[0],'unit':v[1],'aggregation':v[2]} for k,v in METRICS.items()]
    @app.get('/api/accounts/{a}/projects/{p}/sources')
    def sources(a:str,p:str,u=Depends(current)):
        with conn() as c:
            scope(c,u,a,p);return [dict(r) for r in c.execute('SELECT id,metric,mode,created_at FROM traction_sources WHERE account_id=? AND project_id=?',(a,p))]
    @app.post('/api/accounts/{a}/projects/{p}/sources')
    def source(a:str,p:str,b:Source,u=Depends(current)):
        if b.metric not in METRICS or b.mode not in {'spreadsheet','webhook'}:raise HTTPException(422,'Invalid metric or source mode')
        with conn() as c:
            scope(c,u,a,p,True);old=c.execute('SELECT id FROM traction_sources WHERE account_id=? AND project_id=? AND metric=?',(a,p,b.metric)).fetchone();ident=old['id'] if old else uid();secret=secrets.token_hex(32) if b.mode=='webhook' else None
            c.execute('INSERT INTO traction_sources VALUES(?,?,?,?,?,?,?) ON CONFLICT(account_id,project_id,metric) DO UPDATE SET mode=excluded.mode,secret=excluded.secret,created_at=excluded.created_at',(ident,a,p,b.metric,b.mode,secret,now()))
            return {'id':ident,'metric':b.metric,'mode':b.mode,'secret':secret,'endpoint':f'/api/webhooks/{ident}'}
    @app.post('/api/accounts/{a}/projects/{p}/imports/preview')
    async def preview(a:str,p:str,file:UploadFile=File(...),u=Depends(current)):
        with conn() as c:scope(c,u,a,p,True)
        raw=await file.read(2_000_001)
        if len(raw)>2_000_000:raise HTTPException(413,'Maximum upload size is 2 MB')
        try:
            if file.filename.lower().endswith('.csv'):rows=list(csv.DictReader(io.StringIO(raw.decode('utf-8-sig'))))
            elif file.filename.lower().endswith('.xlsx'):
                import zipfile
                with zipfile.ZipFile(io.BytesIO(raw)) as z:
                    if sum(x.file_size for x in z.infolist())>20_000_000:raise ValueError('Workbook too large')
                wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True);it=wb.active.iter_rows(values_only=True);headers=next(it);rows=[]
                for r in it:
                    if len(rows)>10000:raise ValueError('Too many rows')
                    if any(x is not None for x in r):rows.append(dict(zip(headers,r)))
                wb.close()
            else:raise ValueError('Use CSV or XLSX')
        except Exception as e:raise HTTPException(422,'Cannot read spreadsheet: '+str(e))
        rows=valid(rows)
        with conn() as c:
            for metric in {r['metric'] for r in rows}:
                s=c.execute('SELECT mode FROM traction_sources WHERE account_id=? AND project_id=? AND metric=?',(a,p,metric)).fetchone()
                if s and s['mode']!='spreadsheet':raise HTTPException(409,f'{metric} uses a webhook; change its primary source before importing')
            ident=uid();c.execute('INSERT INTO traction_imports VALUES(?,?,?,?,?,0,?)',(ident,a,p,json.dumps(rows),u['id'],now()))
        return {'id':ident,'count':len(rows),'rows':rows[:25],'policy':'Replace metric/date snapshots; never add duplicate values'}
    @app.post('/api/accounts/{a}/projects/{p}/imports/{i}/commit')
    def commit(a:str,p:str,i:str,u=Depends(current)):
        with conn() as c:
            scope(c,u,a,p,True);r=c.execute('SELECT * FROM traction_imports WHERE id=? AND account_id=? AND project_id=? AND actor_id=?',(i,a,p,u['id'])).fetchone()
            if not r:raise HTTPException(404,'Import not found')
            if r['committed']:return {'ok':True,'already_committed':True}
            for row in json.loads(r['rows_json']):
                s=c.execute('SELECT * FROM traction_sources WHERE account_id=? AND project_id=? AND metric=?',(a,p,row['metric'])).fetchone()
                if s and s['mode']!='spreadsheet':raise HTTPException(409,'Source changed; preview import again')
                if not s:
                    ident=uid();c.execute('INSERT INTO traction_sources VALUES(?,?,?,?,?,?,?)',(ident,a,p,row['metric'],'spreadsheet',None,now()))
                else:ident=s['id']
                put(c,a,p,row,ident)
            c.execute('UPDATE traction_imports SET committed=1 WHERE id=?',(i,));return {'ok':True}
    @app.get('/api/accounts/{a}/projects/{p}/traction')
    def traction(a:str,p:str,start:date|None=None,end:date|None=None,u=Depends(current)):
        if start and end and start>end:raise HTTPException(422,'Start must precede end')
        with conn() as c:
            scope(c,u,a,p);rows=[dict(r) for r in c.execute('SELECT v.*,s.mode FROM traction_values v JOIN traction_sources s ON s.id=v.source_id WHERE v.account_id=? AND v.project_id=? AND v.day>=? AND v.day<=? ORDER BY v.day',(a,p,str(start or date.min),str(end or date.max)))]
            metrics=[]
            for k,(label,unit,agg) in METRICS.items():
                rr=[r for r in rows if r['metric']==k];latest=rr[-1] if rr else None
                metrics.append({'key':k,'label':label,'unit':unit,'aggregation':agg,'value':(sum(r['value'] for r in rr) if agg=='flow' else latest['value']) if latest else None,'as_of':latest['day'] if latest else None,'updated_at':latest['updated_at'] if latest else None,'source':latest['mode'] if latest else None})
            return {'metrics':metrics,'history':rows,'timezone':'Africa/Lagos','currency':'NGN','active_definition':'Distinct users performing a meaningful action on the reporting date (DAU)'}
    @app.post('/api/webhooks/{s}')
    async def webhook(s:str,request:Request):
        raw=await request.body()
        if len(raw)>100000:raise HTTPException(413,'Payload too large')
        with conn() as c:
            src=c.execute('SELECT * FROM traction_sources WHERE id=? AND mode="webhook"',(s,)).fetchone()
            if not src:raise HTTPException(404,'Source not found')
            stamp=request.headers.get('x-promptto-timestamp','');sig=request.headers.get('x-promptto-signature','')
            try:
                if abs(time.time()-int(stamp))>300:raise ValueError()
            except ValueError:raise HTTPException(401,'Timestamp outside five-minute window')
            expected=hmac.new(src['secret'].encode(),stamp.encode()+b'.'+raw,hashlib.sha256).hexdigest()
            if not hmac.compare_digest(sig,expected):raise HTTPException(401,'Invalid signature')
            try:
                payload=json.loads(raw);event_id=payload['event_id'];rows=valid([payload])
                if not isinstance(event_id,str) or not 1<=len(event_id)<=160 or rows[0]['metric']!=src['metric']:raise ValueError()
            except (ValueError,KeyError,TypeError):raise HTTPException(422,'Invalid event payload')
            if c.execute('SELECT 1 FROM traction_receipts WHERE source_id=? AND event_id=?',(s,event_id)).fetchone():return {'ok':True,'duplicate':True}
            c.execute('INSERT INTO traction_receipts VALUES(?,?,?)',(s,event_id,now()));put(c,src['account_id'],src['project_id'],rows[0],s)
            return {'ok':True}
